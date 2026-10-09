"""Tests for the cleaning steps and the Excel report.

Run with:  python -m pytest      (or)      python -m unittest discover tests
"""

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from sales_cleaner.__main__ import main
from sales_cleaner.cleaner import clean_sales, load_monthly_files, parse_dates
from sales_cleaner.report import build_report

HEADER = "order_id,order_date,region,product,category,quantity,unit_price\n"


def frame(rows: list[list[str | None]]) -> pd.DataFrame:
    cols = ["order_id", "order_date", "region", "product", "category", "quantity", "unit_price"]
    return pd.DataFrame(rows, columns=cols, dtype="str")


class TestDates(unittest.TestCase):
    def test_mixed_formats(self):
        s = pd.Series(["2026-07-03", "07/04/2026", "5-Jul-2026", "2026/07/06",
                       " 2026-07-07 ", "not a date", None], dtype="str")
        got = parse_dates(s)
        self.assertEqual(
            [d.strftime("%Y-%m-%d") if pd.notna(d) else None for d in got],
            ["2026-07-03", "2026-07-04", "2026-07-05", "2026-07-06", "2026-07-07", None, None])


class TestCleaning(unittest.TestCase):
    def test_removes_exact_and_near_duplicates(self):
        df = frame([
            ["ORD-1", "2026-07-01", "West", "Wireless Mouse", "Accessories", "1", "25"],
            ["ORD-1", "2026-07-01", "West", "Wireless Mouse", "Accessories", "1", "25"],
            ["ORD-1", "07/01/2026", " west", "  wireless mouse ", "Accessories", "1", "$25"],
            ["ORD-2", "2026-07-02", "West", "Wireless Mouse", "Accessories", "2", "25"],
        ])
        clean, log = clean_sales(df)
        self.assertEqual(clean["order_id"].tolist(), ["ORD-1", "ORD-2"])
        self.assertEqual(log.exact_duplicates_removed, 2)

    def test_keeps_first_when_order_id_repeats_with_different_values(self):
        df = frame([
            ["ORD-1", "2026-07-01", "West", "USB-C Hub", "Accessories", "1", "40"],
            ["ORD-1", "2026-07-01", "West", "USB-C Hub", "Accessories", "3", "40"],
        ])
        clean, log = clean_sales(df)
        self.assertEqual(len(clean), 1)
        self.assertEqual(clean.loc[0, "quantity"], 1)
        self.assertEqual(log.repeated_order_ids_removed, 1)

    def test_canonical_spelling_and_whitespace(self):
        df = frame([
            ["ord-1", "2026-07-01", "West", "USB-C Hub", "Accessories", "1", "40"],
            ["ORD-2", "2026-07-01", "WEST", "usb-c   hub", "Accessories", "1", "40"],
            ["ORD-3", "2026-07-01", "West", "USB-C Hub", "Accessories", "1", "40"],
        ])
        clean, _ = clean_sales(df)
        self.assertEqual(set(clean["product"]), {"USB-C Hub"})
        self.assertEqual(set(clean["region"]), {"West"})
        self.assertIn("ORD-1", clean["order_id"].tolist())

    def test_fills_missing_values(self):
        df = frame([
            ["ORD-1", "2026-07-01", "West", "Mouse", "Accessories", "2", "20"],
            ["ORD-2", "2026-07-02", "East", "Mouse", "Accessories", "4", "30"],
            ["ORD-3", "2026-07-03", None, "Mouse", None, None, None],
        ])
        clean, log = clean_sales(df)
        row = clean.set_index("order_id").loc["ORD-3"]
        self.assertEqual(row["unit_price"], 25.0)     # median of 20 and 30
        self.assertEqual(row["quantity"], 3)          # median of 2 and 4
        self.assertEqual(row["region"], "Unknown")
        self.assertEqual(row["category"], "Accessories")
        self.assertEqual(row["revenue"], 75.0)
        self.assertEqual(log.values_filled, 4)
        self.assertFalse(clean.isna().any().any())

    def test_drops_rows_that_cannot_be_used(self):
        df = frame([
            ["ORD-1", "sometime", "West", "Mouse", "Accessories", "1", "20"],
            [None, "2026-07-01", "West", "Mouse", "Accessories", "1", "20"],
            ["ORD-3", "2026-07-01", "West", None, "Accessories", "1", "20"],
            ["ORD-4", "2026-07-01", "West", "Mouse", "Accessories", "1", "20"],
        ])
        clean, log = clean_sales(df)
        self.assertEqual(clean["order_id"].tolist(), ["ORD-4"])
        self.assertEqual(log.rows_dropped_bad_date, 1)
        self.assertEqual(log.rows_dropped_missing_key, 2)

    def test_prices_with_symbols_become_numbers(self):
        df = frame([["ORD-1", "2026-07-01", "West", "Monitor", "Displays", "2", "$1,299.50"]])
        clean, _ = clean_sales(df)
        self.assertEqual(clean.loc[0, "unit_price"], 1299.5)
        self.assertEqual(clean.loc[0, "revenue"], 2599.0)
        self.assertEqual(clean.loc[0, "month"], "2026-07")


class TestEndToEnd(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        raw = self.root / "raw"
        raw.mkdir()
        (raw / "sales_2026_07.csv").write_text(HEADER +
            "ORD-1,2026-07-01,West,Mouse,Accessories,1,20\n"
            "ORD-1,2026-07-01,West,Mouse,Accessories,1,20\n"
            "ORD-2,07/15/2026,East,Monitor,Displays,1,300\n")
        (raw / "sales_2026_08.csv").write_text(HEADER +
            "ORD-3,2026-08-02,West,Mouse,Accessories,5,$20\n"
            "ORD-4,3-Aug-2026,,Monitor,Displays,,300\n")

    def tearDown(self):
        self.tmp.cleanup()

    def test_loader_tracks_source_file(self):
        df, files = load_monthly_files(self.root / "raw")
        self.assertEqual(files, ["sales_2026_07.csv", "sales_2026_08.csv"])
        self.assertEqual(len(df), 5)
        self.assertEqual(df["source_file"].nunique(), 2)

    def test_missing_folder_or_columns(self):
        with self.assertRaises(FileNotFoundError):
            load_monthly_files(self.root / "nope")
        bad = self.root / "bad"
        bad.mkdir()
        (bad / "x.csv").write_text("order_id,amount\n1,2\n")
        with self.assertRaises(ValueError):
            load_monthly_files(bad)

    def test_cli_writes_csv_and_report(self):
        out = self.root / "out"
        code = main(["-i", str(self.root / "raw"), "-o", str(out)])
        self.assertEqual(code, 0)

        clean = pd.read_csv(out / "clean_sales.csv")
        self.assertEqual(clean["order_id"].tolist(), ["ORD-1", "ORD-2", "ORD-3", "ORD-4"])
        self.assertEqual(clean["order_date"].tolist(),
                         ["2026-07-01", "2026-07-15", "2026-08-02", "2026-08-03"])

        wb = load_workbook(out / "sales_report.xlsx")
        self.assertEqual(wb.sheetnames, ["Summary", "Clean Data", "Data Quality"])
        summary = wb["Summary"]
        labels = [summary.cell(row=5, column=c).value for c in range(2, 7)]
        self.assertEqual(labels, ["Total Revenue", "Orders", "Units Sold",
                                  "Average Order Value", "Top Product (Revenue)"])
        # Headline figures are live formulas, not pasted numbers
        for c in range(2, 7):
            self.assertTrue(str(summary.cell(row=6, column=c).value).startswith("="))
        self.assertEqual(len(summary._charts), 1)
        self.assertEqual(wb["Clean Data"].max_row, 5)  # header + 4 rows

    def test_report_rejects_empty_data(self):
        df = frame([["ORD-1", "bad date", "West", "Mouse", "Accessories", "1", "20"]])
        clean, log = clean_sales(df)
        with self.assertRaises(ValueError):
            build_report(clean, log, self.root / "empty.xlsx")


if __name__ == "__main__":
    unittest.main()
