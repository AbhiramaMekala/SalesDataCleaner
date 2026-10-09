"""Build a formatted Excel report from the cleaned sales data.

Sheets
------
Summary       5 headline figures, revenue by month (with a chart),
              revenue by region and by product
Clean Data    every cleaned row, formatted, filterable
Data Quality  what the cleaner changed (duplicates removed, values filled ...)

Every number on the Summary sheet is an Excel formula that points at the
Clean Data sheet (SUM, SUMIFS, COUNTIFS, INDEX/MATCH), not a pasted value,
so the report stays correct if someone edits or filters the data.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .cleaner import CleaningLog

FONT = "Arial"
NAVY = "1F3864"
LIGHT = "EEF2F8"
GREY = "7F7F7F"
MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
INT = "#,##0"
PCT = "0.0%"

THIN = Side(style="thin", color="D0D7E2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DATA_SHEET = "Clean Data"
DATA = f"'{DATA_SHEET}'"  # quoted because the name has a space

# Column letters on the Clean Data sheet (must match OUTPUT_COLUMNS order)
COL = {"order_id": "A", "order_date": "B", "month": "C", "region": "D", "product": "E",
       "category": "F", "quantity": "G", "unit_price": "H", "revenue": "I", "source_file": "J"}


def _font(**kw) -> Font:
    return Font(name=FONT, **kw)


def _rng(col: str, last_row: int) -> str:
    c = COL[col]
    return f"{DATA}!${c}$2:${c}${last_row}"


def _header_row(ws, row: int, start_col: int, labels: list[str]) -> None:
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + i, value=label)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX


def _section_title(ws, row: int, text: str) -> None:
    cell = ws.cell(row=row, column=2, value=text)
    cell.font = _font(bold=True, size=12, color=NAVY)


# ---------------------------------------------------------------- sheets

def _write_data_sheet(wb: Workbook, df: pd.DataFrame) -> int:
    ws = wb.create_sheet(DATA_SHEET)
    headers = ["Order ID", "Order Date", "Month", "Region", "Product", "Category",
               "Quantity", "Unit Price", "Revenue", "Source File"]
    ws.append(headers)
    for rec in df.itertuples(index=False):
        ws.append([rec.order_id, rec.order_date.to_pydatetime().date(), rec.month,
                   rec.region, rec.product, rec.category, int(rec.quantity),
                   float(rec.unit_price), float(rec.revenue), rec.source_file])
    last = ws.max_row

    for cell in ws[1]:
        cell.font = _font(bold=True, color="FFFFFF")
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = _font()
        row[1].number_format = "yyyy-mm-dd"
        row[6].number_format = INT
        row[7].number_format = MONEY
        row[8].number_format = MONEY

    table = Table(displayName="CleanSales", ref=f"A1:J{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    for col, width in zip("ABCDEFGHIJ", [13, 12, 10, 12, 28, 13, 10, 12, 13, 20]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    return last


def _write_quality_sheet(wb: Workbook, log: CleaningLog) -> None:
    ws = wb.create_sheet("Data Quality")
    ws["B2"] = "What the cleaner changed"
    ws["B2"].font = _font(bold=True, size=14, color=NAVY)
    ws["B3"] = "Counts recorded by the cleaning script during this run."
    ws["B3"].font = _font(italic=True, color=GREY)

    _header_row(ws, 5, 2, ["Step", "Count"])
    r = 6
    for label, value in log.as_rows():
        ws.cell(row=r, column=2, value=label).font = _font()
        c = ws.cell(row=r, column=3, value=value)
        c.font = _font(bold=label in ("Rows read", "Rows written"))
        c.number_format = INT
        for col in (2, 3):
            ws.cell(row=r, column=col).border = BOX
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="Source files").font = _font(bold=True)
    for name in log.files_read:
        r += 1
        ws.cell(row=r, column=2, value=name).font = _font()
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 12


def _write_summary_sheet(ws, df: pd.DataFrame, last: int, log: CleaningLog) -> None:
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    for col in "CDEF":
        ws.column_dimensions[col].width = 20
    ws.column_dimensions["G"].width = 3
    # Print / PDF: landscape, everything one page wide
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    months = sorted(df["month"].unique())
    period = (f"{_month_label(months[0])} – {_month_label(months[-1])}"
              if len(months) > 1 else _month_label(months[0]))

    ws["B2"] = "Sales Report"
    ws["B2"].font = _font(bold=True, size=18, color=NAVY)
    ws["B3"] = (f"{period}  ·  {len(log.files_read)} files  ·  "
                f"generated {datetime.now():%Y-%m-%d %H:%M}")
    ws["B3"].font = _font(color=GREY)

    # ---- 5 headline figures -------------------------------------------
    rev, qty, oid = _rng("revenue", last), _rng("quantity", last), _rng("order_id", last)
    kpis = [
        ("Total Revenue", f"=SUM({rev})", MONEY0),
        ("Orders", f"=COUNTA({oid})", INT),
        ("Units Sold", f"=SUM({qty})", INT),
        ("Average Order Value", "=IFERROR(B6/C6,0)", MONEY),
        ("Top Product (Revenue)", None, None),  # filled once the product table exists
    ]
    for i, (label, formula, fmt) in enumerate(kpis):
        col = 2 + i
        head = ws.cell(row=5, column=col, value=label)
        head.font = _font(bold=True, size=9, color=GREY)
        head.fill = PatternFill("solid", fgColor=LIGHT)
        head.alignment = Alignment(horizontal="center")
        val = ws.cell(row=6, column=col, value=formula)
        val.font = _font(bold=True, size=16 if i < 4 else 11, color=NAVY)
        val.fill = PatternFill("solid", fgColor=LIGHT)
        val.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if fmt:
            val.number_format = fmt
    ws.row_dimensions[6].height = 34

    # ---- Revenue by month (+ chart) -----------------------------------
    row = 9
    _section_title(ws, row, "Revenue by Month")
    _header_row(ws, row + 1, 2, ["Month", "Orders", "Units", "Revenue", "Share of Revenue"])
    first_month_row = row + 2
    for i, m in enumerate(months):
        r = first_month_row + i
        ws.cell(row=r, column=2, value=m)  # text "YYYY-MM", matches the Month column
        # SUMPRODUCT compares text exactly; COUNTIFS/SUMIFS would try to read
        # "2026-07" as a date and miss the text values in the Month column.
        is_month = f'--({_rng("month", last)}=B{r})'
        ws.cell(row=r, column=3, value=f"=SUMPRODUCT({is_month})").number_format = INT
        ws.cell(row=r, column=4, value=f"=SUMPRODUCT({is_month},{qty})").number_format = INT
        ws.cell(row=r, column=5, value=f"=SUMPRODUCT({is_month},{rev})").number_format = MONEY0
        ws.cell(row=r, column=6, value=f"=IFERROR(E{r}/$B$6,0)").number_format = PCT
    last_month_row = first_month_row + len(months) - 1
    total_r = last_month_row + 1
    _total_row(ws, total_r, first_month_row, last_month_row)
    _style_body(ws, first_month_row, total_r, 2, 6)

    chart = BarChart()
    chart.type = "col"
    chart.title = "Monthly Revenue"
    chart.y_axis.title = "Revenue ($)"
    chart.y_axis.numFmt = '"$"#,##0'
    chart.y_axis.majorGridlines = None
    chart.legend = None
    chart.add_data(Reference(ws, min_col=5, min_row=row + 1, max_row=last_month_row),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=first_month_row, max_row=last_month_row))
    chart.series[0].graphicalProperties.solidFill = NAVY
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showSerName = False
    chart.dataLabels.showCatName = False
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.numFmt = '"$"#,##0'
    chart.height, chart.width = 7.5, 15
    ws.add_chart(chart, "H9")

    # ---- Revenue by region --------------------------------------------
    row = total_r + 3
    _section_title(ws, row, "Revenue by Region")
    _header_row(ws, row + 1, 2, ["Region", "Orders", "Units", "Revenue", "Share of Revenue"])
    regions = (df.groupby("region")["revenue"].sum().sort_values(ascending=False).index.tolist())
    start = row + 2
    for i, name in enumerate(regions):
        r = start + i
        ws.cell(row=r, column=2, value=name)
        _group_formulas(ws, r, "region", last)
    end = start + len(regions) - 1
    _total_row(ws, end + 1, start, end)
    _style_body(ws, start, end + 1, 2, 6)

    # ---- Revenue by product -------------------------------------------
    row = end + 4
    _section_title(ws, row, "Revenue by Product")
    _header_row(ws, row + 1, 2, ["Product", "Orders", "Units", "Revenue", "Share of Revenue"])
    products = (df.groupby("product")["revenue"].sum().sort_values(ascending=False).index.tolist())
    p_start = row + 2
    for i, name in enumerate(products):
        r = p_start + i
        ws.cell(row=r, column=2, value=name)
        _group_formulas(ws, r, "product", last)
    p_end = p_start + len(products) - 1
    _total_row(ws, p_end + 1, p_start, p_end)
    _style_body(ws, p_start, p_end + 1, 2, 6)

    # Top product = the product with the highest revenue in the table above
    ws["F6"] = (f"=INDEX($B${p_start}:$B${p_end},"
                f"MATCH(MAX($E${p_start}:$E${p_end}),$E${p_start}:$E${p_end},0))")

    note_r = p_end + 3
    ws.cell(row=note_r, column=2,
            value="All figures are formulas over the 'Clean Data' sheet. "
                  "See 'Data Quality' for what was cleaned.").font = _font(italic=True, size=9, color=GREY)


def _group_formulas(ws, r: int, col: str, last: int) -> None:
    key = _rng(col, last)
    ws.cell(row=r, column=3, value=f"=COUNTIFS({key},B{r})").number_format = INT
    ws.cell(row=r, column=4, value=f'=SUMIFS({_rng("quantity", last)},{key},B{r})').number_format = INT
    ws.cell(row=r, column=5, value=f'=SUMIFS({_rng("revenue", last)},{key},B{r})').number_format = MONEY0
    ws.cell(row=r, column=6, value=f"=IFERROR(E{r}/$B$6,0)").number_format = PCT


def _total_row(ws, r: int, first: int, last: int) -> None:
    ws.cell(row=r, column=2, value="Total")
    for col, fmt in ((3, INT), (4, INT), (5, MONEY0), (6, PCT)):
        letter = get_column_letter(col)
        c = ws.cell(row=r, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
        c.number_format = fmt
    for col in range(2, 7):
        ws.cell(row=r, column=col).font = _font(bold=True)
        ws.cell(row=r, column=col).border = Border(top=Side(style="medium", color=NAVY))


def _style_body(ws, first: int, last: int, c1: int, c2: int) -> None:
    for r in range(first, last + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if r < last:
                cell.font = _font()
                cell.border = BOX
                if (r - first) % 2:
                    cell.fill = PatternFill("solid", fgColor="F7F9FC")
            cell.alignment = Alignment(horizontal="left" if c == c1 else "right")


def _month_label(m: str) -> str:
    return datetime.strptime(m, "%Y-%m").strftime("%b %Y")


# ---------------------------------------------------------------- entry point

def build_report(df: pd.DataFrame, log: CleaningLog, path: str | Path) -> Path:
    if df.empty:
        raise ValueError("No rows left after cleaning; nothing to report.")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    summary = wb.active
    last = _write_data_sheet(wb, df)
    _write_summary_sheet(summary, df, last, log)
    _write_quality_sheet(wb, log)

    wb.calculation.fullCalcOnLoad = True  # Excel computes every formula on open
    wb.save(path)
    return path
